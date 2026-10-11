"""Freeze child labels and audit actual contract/data. Original caches untouched."""
import json
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from openpyxl import load_workbook
from training.labels import sha256
from weld_pipeline.curved_cache import load_piece
from weld_pipeline.contract_child_evaluation import child_units, score_children, aggregate

ROOT = Path('weld_pipeline/data/contract_child_v1')
REPORT = Path('weld_pipeline/reports/contract_child_20260930')
ATTRIBUTES = ['thickness_1','thickness_2','normal_1','normal_2','front_wall','back_wall',
    'start_height','end_height','start_hole_height','end_hole_height','start_base_edge','end_base_edge',
    'has_eave','side_plate_split']
FIELDS = {'front_wall':'FrontWallDistance','back_wall':'BackWallDistance',
    'start_height':'startHeight','end_height':'endHeight','start_hole_height':'StartViaHoleHeight',
    'end_hole_height':'EndViaHoleHeight','start_base_edge':'StartBasePlateEdgeDistance',
    'end_base_edge':'EndBasePlateEdgeDistance','has_eave':'HasEave (Any)',
    'side_plate_split':'IsSplitedForDualRobot (Any)'}


def child_attributes(parent, index):
    meta = parent.get('metadata', {})
    last = len(parent['segments']) - 1
    out = {}
    for name in ATTRIBUTES:
        value = meta.get(FIELDS.get(name, ''))
        reason = 'missing_customer_truth'
        available = value is not None
        if name.startswith(('start_', 'front_')) and index != 0:
            available, reason = False, 'internal_child_endpoint_not_labelled'
        elif name.startswith(('end_', 'back_')) and index != last:
            available, reason = False, 'internal_child_endpoint_not_labelled'
        elif name in ['has_eave','side_plate_split'] and last:
            available, reason = False, 'parent_any_flag_not_child_truth'
        if name in ['front_wall','back_wall'] and value == -1:
            available, reason = False, 'legacy_no_object_or_out_of_range_ambiguous'
        state = 'value'
        if available and name.endswith('base_edge') and value == 51:
            state = 'greater_than_50_mm'
        if available and name.endswith('hole_height') and value == 0:
            state = 'no_hole'
        out[name] = {'value': value if available else None, 'state': state if available else 'missing_truth',
            'truth_available': available, 'reason': None if available else reason, 'source': 'customer_excel' if available else None}
    return out


def main():
    ROOT.mkdir(exist_ok=True, parents=True); REPORT.mkdir(exist_ok=True, parents=True)
    source = Path('weld_pipeline/data/partial_augmented_surface_v4')
    split = json.loads((source/'split.json').read_text('utf8'))
    contract = Path(r'C:\Users\Administrator\Desktop\电动舵机力矩在线辨识\点云焊缝提取技术开发合同_修改.docx')
    xml = ET.fromstring(zipfile.ZipFile(contract).read('word/document.xml'))
    ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    text = '\n'.join(''.join(t.text or '' for t in p.findall('.//w:t',ns)) for p in xml.findall('.//w:p',ns))
    (REPORT/'contract_text.txt').write_text(text, encoding='utf8')
    header_records = []
    for path in sorted(Path('DataSet').rglob('*.xlsx')):
        book=load_workbook(path,read_only=True,data_only=True)
        headers=[str(v or '') for v in next(book.active.iter_rows(values_only=True))]; book.close()
        header_records.append({'path':str(path),'sha256':sha256(path),'columns':len(headers),
            'thickness_normal_headers':[h for h in headers if any(k in h.lower() for k in ['thickness','normal','板厚','法向'])]})
    all_rows, scores, records = [], {}, []
    manifests = {}
    independent = [n for n in split['train'] if '__' not in n]
    for name in independent + split['validation']:
        folder=(source/split['piece_roots'][name]).resolve()
        piece=load_piece(folder, verify=True)
        if sha256(piece['record']['input']) != piece['record']['pcd_sha256']:
            raise ValueError('Raw PCD changed: '+name)
        for path,wanted in piece['record']['reference_hashes'].items():
            if sha256(path)!=wanted: raise ValueError('Reference changed: '+path)
        children = child_units(piece['truth'])
        for child in children:
            child['attributes'] = child_attributes(piece['truth'][child['parent_index']], child['segment_index'])
            child['workpiece_id'] = name
        (ROOT/(name+'_children.json')).write_text(json.dumps(children,ensure_ascii=False,indent=2),encoding='utf8')
        all_rows.extend(children)
        role='train' if name in independent else 'development'
        records.append({'name':name,'role':role,'parents':len(piece['truth']),'children':len(children),
            'pcd':piece['record']['input'],'pcd_sha256':piece['record']['pcd_sha256'],
            'reference_hashes':piece['record']['reference_hashes']})
        manifests[name]=sha256(folder/'manifest.json')
        if role=='development':
            prediction_path=Path('weld_pipeline/runs/typed_branch_surface_v3_seed115/final_evaluation')/(name+'_predictions.json')
            pred=json.loads(prediction_path.read_text('utf8'))
            scores[name]=score_children(pred,piece['truth'])
            (REPORT/(name+'_baseline_score.json')).write_text(json.dumps(scores[name],indent=2),encoding='utf8')
    copied={**split,'status':'FROZEN_CONTRACT_CHILD_DEVELOPMENT_V1',
        'source_split_sha256':sha256(source/'split.json'), 'count_unit':'each existing ordered geometric segment',
        'mapping_policy':'Provisional explicit mapping from customer child-count answer; preserve all existing real breakpoints, no equivalence regrouping.',
        'piece_roots':{name:str((source/path).resolve().relative_to(Path.cwd())) for name,path in split['piece_roots'].items()},
        'labels_only_migration':True,'runtime_step_forbidden':True,'model_input':'AllFaces XYZ only',
        'timing_gate':'DEFERRED_BY_USER_NOT_PASSED'}
    (ROOT/'split.json').write_text(json.dumps(copied,ensure_ascii=False,indent=2),encoding='utf8')
    attribute_counts={a:{'available':sum(c['attributes'][a]['truth_available'] for c in all_rows),
        'missing':sum(not c['attributes'][a]['truth_available'] for c in all_rows),
        'acceptance_status':'NOT_RUN_MODEL_ATTRIBUTE_OUTPUTS_NOT_IMPLEMENTED'} for a in ATTRIBUTES}
    audit={'contract_sha256':sha256(contract),'excel_headers':header_records,'independent_workpieces':records,
        'train_workpieces':len(independent),'dev_workpieces':len(split['validation']),
        'attribute_truth_coverage':attribute_counts,'cached_manifest_hashes':manifests,
        'baseline_dev':aggregate(scores),'source_checkpoint_sha256':sha256('weld_pipeline/runs/typed_branch_surface_v3_seed115/best.pt'),
        'score_source_sha256':sha256('weld_pipeline/contract_child_evaluation.py'),
        'full_contract_status':'NOT_ACHIEVED','missing_side_truth':True,'independent_test_available':False,
        'no_independent_curve_validation':True,'timing':'DEFERRED_BY_USER'}
    (REPORT/'audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({k:v for k,v in audit.items() if k in ['train_workpieces','dev_workpieces','attribute_truth_coverage','baseline_dev','full_contract_status']},ensure_ascii=False),flush=True)


if __name__=='__main__': main()
