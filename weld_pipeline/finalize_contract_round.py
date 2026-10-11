"""Verify sources/exports and freeze an honest summary of this development round."""
import csv
import json
from pathlib import Path
from openpyxl import load_workbook
from plyfile import PlyData
from training.labels import sha256


def main():
    report=Path('weld_pipeline/reports/contract_child_20260930')
    original=json.loads((report/'audit.json').read_text('utf8'))
    for item in original['excel_headers']:
        if sha256(item['path'])!=item['sha256']:raise ValueError('Original workbook changed')
    for item in original['independent_workpieces']:
        if sha256(item['pcd'])!=item['pcd_sha256']:raise ValueError('Original PCD changed')
        for path,wanted in item['reference_hashes'].items():
            if sha256(path)!=wanted:raise ValueError('Original truth changed')
    if sha256('weld_pipeline/runs/typed_branch_surface_v3_seed115/best.pt')!=original['source_checkpoint_sha256']:
        raise ValueError('Retained baseline changed')
    runs=[]
    for name in ['contract_extent_v1_seed121','contract_horizontal_v1_seed122','contract_context_v1_seed123']:
        folder=Path('weld_pipeline/runs')/name
        complete=json.loads((folder/'complete.json').read_text('utf8'))
        audit=json.loads((folder/'contract_goal_audit.json').read_text('utf8'))
        rows=[json.loads(s) for s in (folder/'epochs.jsonl').read_text('utf8').splitlines() if s.strip()]
        if [r['epoch'] for r in rows]!=list(range(1,121)):raise ValueError('Incomplete training epochs')
        if complete['epochs']!=120:raise ValueError('Incomplete training')
        runs.append({'run':name,'epochs':120,'selected_epoch':audit['selected_epoch'],
            'best_children':audit['development']['valid_children'],'truth_children':audit['development']['truth_count'],
            'false_children':audit['development']['false_predictions'],
            'last_epoch_children':rows[-1]['summary']['valid_children'],
            'elapsed_training_seconds':complete['elapsed_seconds'],'status':audit['status'],
            'checkpoint_sha256':sha256(folder/'best.pt'),'audit_sha256':sha256(folder/'contract_goal_audit.json'),
            'every_epoch_frozen_branch_verified':all(r.get('horizontal_unchanged',r.get('vertical_unchanged',False)) for r in rows)})
    base=Path('weld_pipeline/runs/contract_horizontal_v1_seed122')
    exports=[]
    for folder in sorted(base.glob('raw_*')):
        pred=json.loads((folder/'predictions.json').read_text('utf8'))
        receipt=json.loads((folder/'inference_receipt.json').read_text('utf8'))
        with (folder/'predicted_children.csv').open(encoding='utf8',newline='') as f:
            child_rows=list(csv.DictReader(f))
        expected=sum(len(p['segments']) for p in pred)
        if len(child_rows)!=expected or expected!=receipt['child_predictions']:raise ValueError('Child export count mismatch')
        curves=json.loads((folder/'child_curves.json').read_text('utf8'))
        for row in child_rows:
            curve=curves[row['weld_id']]
            if not curve['ordered_points_mm']:raise ValueError('Missing curve reference')
            for field in ['thickness_1','thickness_2','normal_1','normal_2']:
                if row[field+'_value']!='' or row[field+'_status']!='unknown':raise ValueError('Fabricated attribute value')
        with (folder/'predicted_welds.csv').open(encoding='utf-8-sig',newline='') as f:
            parent_rows=list(csv.DictReader(f))
        if len(parent_rows)!=len(pred):raise ValueError('Parent CSV count mismatch')
        book=load_workbook(folder/'预测焊缝.xlsx',read_only=True,data_only=False)
        if book['焊缝汇总'].max_row-1!=len(pred):raise ValueError('Workbook count mismatch')
        book.close()
        with (folder/'预测焊缝.ply').open('rb') as f:
            ply=PlyData.read(f)
        if pred and not len(ply['vertex'].data):raise ValueError('Empty prediction layer')
        exports.append({'workpiece':folder.name,'child_rows':expected,'parents':len(pred),
            'raw_receipt_sha256':sha256(folder/'inference_receipt.json'),'xlsx_csv_ply_parse_and_count':'PASS',
            'unknown_attributes_explicit':'PASS'})
    template=Path('合同/客户补充真值_子段样表_20260930.csv')
    with template.open(encoding='utf-8-sig',newline='') as f:review=list(csv.DictReader(f))
    if len(review)!=1605:raise ValueError('Review template count mismatch')
    if any(r['thickness_1_value'] or r['normal_1_x'] or r['normal_2_z'] for r in review):
        raise ValueError('New truth was fabricated')
    result={'round_status':'NOT_ACHIEVED','data_origin':'CAD sampled points',
        'original_pcd_reference_hashes_unchanged':25,'original_main_workbook_hashes_unchanged':22,
        'baseline_115_unchanged':True,'new_scope':'Development child geometry only; existing geometric-segment mapping pending review.',
        'runs':runs,'raw_exports_verified':exports,'customer_review_rows':len(review),
        'customer_review_template_sha256':sha256(template),'timing_gate':'DEFERRED_BY_USER_NOT_PASSED',
        'attributes':'MISSING_REVIEWED_THICKNESS_NORMAL_TRUTH_AND_COMPLETE_CHILD_ATTRIBUTE_TRUTH',
        'segmentation_question':'GR120mm has 3 marked children whereas FR120mm has 1; general rule required.',
        'next_step':'Obtain segmentation rule and reviewed attributes; add independent real single-vertical/curve workpieces, preserve sealed test split.',
        'formal_independent_test':'NOT_AVAILABLE','new_models_promoted':False}
    (report/'round_completion_receipt.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({'status':result['round_status'],'verified_training_runs':len(runs),
        'verified_raw_exports':len(exports),'review_rows':len(review),'baseline_preserved':True},ensure_ascii=False),flush=True)


if __name__=='__main__':main()
