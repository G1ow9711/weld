"""Read original bilingual headers and audit process-label censoring."""
import json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook
from training.labels import sha256
from weld_pipeline.truth import excel_rows
from weld_pipeline.process_label_semantics import process_labels,WALL_FIELDS,EDGE_FIELDS,HEIGHT_FIELDS,SEMANTICS_VERSION


def main():
    out=Path('weld_pipeline/reports/process_semantics');out.mkdir(parents=True,exist_ok=True)
    records=[];sources=[];headers=[]
    for path in sorted(Path('DataSet/9月20日').glob('*/*.xlsx')):
        book=load_workbook(path,read_only=True,data_only=True)
        hs=list(next(book.active.iter_rows(values_only=True)));book.close()
        headers.append({'file':str(path),'headers':hs})
        sources.append({'file':str(path),'sha256':sha256(path)})
        for row in excel_rows(path):
            records.append({'group':path.parent.name,'id':str(row['ObjectId']),
                'number':row['WeldNumber (f/v + Name number)'],**process_labels(row,path.parent.name)})
    stats={}
    for field in WALL_FIELDS+EDGE_FIELDS+HEIGHT_FIELDS:
        values=[r['distances'][field] for r in records]
        exact=[r['exact_mm'] for r in values if r['exact_mm'] is not None]
        stats[field]={'states':dict(Counter(r['state'] for r in values)),
                      'raw_top_values':Counter(str(r['raw']) for r in values).most_common(8),
                      'exact_range_mm':[min(exact),max(exact)] if exact else None}
    report={'scope':'Original 22-group Excel audit; semantic correction from customer-provided diagrams.',
        'semantics_version':SEMANTICS_VERSION,
        'customer_confirmed':{'distance_unit':'millimeter','edge_distance_over_50_encoding':51,'edge_distance_50_is_exact':True,
            'passage_requires_split_when':'distance_mm <= 35','zero_via_hole_height':'no_hole',
            'geometric_segments_and_breakpoints_required':True,
            '224-FR71A_f1_f2_EndBasePlateEdgeDistance':51},
        'groups':len(sources),'welds':len(records),'sources':sources,'headers':headers,'distance_statistics':stats,
        'side_plate_split_count':sum(r['side_plate_split_any'] is True for r in records),
        'intersection_distance_columns':[h for item in headers for h in item['headers'] if h and 'IntersectionDistance' in str(h)],
        'label_corrections':[c for r in records for c in r['label_corrections']],
        'pending':['Other negative/zero plate-height sentinels, when present; no additional meaning inferred.']}
    (out/'audit.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
    with (out/'supervision.jsonl').open('w',encoding='utf-8') as f:
        for row in records:f.write(json.dumps(row,ensure_ascii=False)+'\n')
    print(json.dumps({k:report[k] for k in ['groups','welds','side_plate_split_count','intersection_distance_columns','distance_statistics']},ensure_ascii=False,indent=2))


if __name__=='__main__':main()
