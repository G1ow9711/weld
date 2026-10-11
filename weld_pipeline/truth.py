"""Customer reference loading for TRAINING/EVALUATION ONLY. No inference imports."""
import json
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from weld_pipeline.evaluation import welding_type
from weld_pipeline.customer_label_corrections import apply_label_corrections


def excel_rows(path):
    book=load_workbook(path,read_only=True,data_only=True)
    iterator=book.active.iter_rows(values_only=True)
    headers=next(iterator)
    names=[str(h).split('\n')[-1] if h is not None else '' for h in headers]
    out=[]
    for row in iterator:
        r=dict(zip(names,row))
        if r.get('ObjectId'):out.append(r)
    book.close()
    return out


def load_truth(folder):
    folder=Path(folder)
    rows=excel_rows(next(folder.glob('*.xlsx')))
    rows=[apply_label_corrections(folder.name,row)[0] for row in rows]
    by_id={str(r['ObjectId']):r for r in rows}
    paths=list(folder.glob('*welds.json'))
    welds=[];excluded=[]
    if paths:
        source=json.loads(paths[0].read_text(encoding='utf-8-sig'))
        if source['unit']!='millimeter' or source['matrixConvention']!='row_major_column_vector':
            raise ValueError('Unsupported curve units/matrix convention')
        if source['coordinateFrame']!='reference_workpiece':raise ValueError('Unsupported reference frame')
        matrix=np.asarray(source['exportToWorld'],dtype=float)
        if not np.allclose(matrix[3],[0,0,0,1]) or not np.allclose(matrix[:3,:3].T@matrix[:3,:3],np.eye(3),atol=1e-7):
            raise ValueError('Non-rigid curve transform')
        for w in source['welds']:
            r=by_id[str(w['id'])]
            out={'id':w['id'],'name':w['name'],'type':welding_type(r['WeldType (isVerticalWeldLine)']),
                 'segments':[],'metadata':r,'source':'exported_curve_json'}
            for s in w['segments']:
                p=np.asarray(s['points'],dtype=float)
                if p.ndim!=2 or p.shape[1]!=3 or len(p)<2 or not np.isfinite(p).all():raise ValueError('Invalid curve samples')
                p=p@matrix[:3,:3].T+matrix[:3,3]
                if s['kind'] not in ('line','arc'):raise ValueError('Unsupported curve primitive')
                out['segments'].append({'kind':s['kind'],'points':p.tolist(),'length_mm':s['lengthMm']})
            welds.append(out)
    else:
        for r in rows:
            wid=str(r['ObjectId'])
            if r['SegmentCount']!=1 or r['HasMid (Any)']:
                excluded.append({'id':wid,'reason':'Only global Excel endpoints; internal segment/arc geometry missing'})
                continue
            a=[float(r['Source StartMatrix.Translation.'+d]) for d in 'XYZ']
            b=[float(r['Source EndMatrix.Translation.'+d]) for d in 'XYZ']
            welds.append({'id':wid,'name':r['Name'],'type':welding_type(r['WeldType (isVerticalWeldLine)']),
                'segments':[{'kind':'line','points':[a,b],'length_mm':float(r['Length'])}],
                'metadata':r,'source':'single_straight_excel_endpoints'})
    return welds,excluded


def sample_weld(weld,step_mm=2.):
    chunks=[]
    for s in weld['segments']:
        p=np.asarray(s['points'],dtype=float)
        if s['kind']=='line':
            length=np.linalg.norm(p[-1]-p[0]);p=np.linspace(p[0],p[-1],max(2,int(np.ceil(length/step_mm))+1))
        chunks.append(p)
    return np.concatenate(chunks)
