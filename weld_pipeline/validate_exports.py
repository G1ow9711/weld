"""Read back a prediction export and compare all coordinates and segment IDs."""
import argparse
import csv
import json
import zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
import numpy as np
from openpyxl import load_workbook
from plyfile import PlyData
from weld_pipeline.export import curve_points


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--outdir',required=True)
    args=ap.parse_args()
    out=Path(args.outdir)
    predictions=json.loads((out/'predictions.json').read_text(encoding='utf-8'))
    expected=[]
    for wi,w in enumerate(predictions,1):
        expected.extend((wi,w['id'],si,pi+1,kind,q,d) for si,pi,kind,q,d in curve_points(w))
    # Keep formula text when reading the workbook. Inspect cached numeric XML
    # separately and verify it by independent geometry calculations below.
    namespace={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(out/'预测焊缝.xlsx') as archive:
        cached=[]
        for name in ['sheet1.xml','sheet2.xml']:
            root=ET.fromstring(archive.read('xl/worksheets/'+name))
            cached.append({c.attrib['r']:c.find('m:v',namespace).text for c in root.findall('.//m:c',namespace)
                           if c.find('m:v',namespace) is not None})
    book=load_workbook(out/'预测焊缝.xlsx',read_only=True,data_only=False)
    summary=list(book['焊缝汇总'].values)[1:]
    assert len(summary)==len(predictions)
    for row_index,(row,w) in enumerate(zip(summary,predictions),2):
        assert row[0]==w['id']
        assert row[1]==('平焊' if w['type']=='horizontal' else '立焊')
        ends=np.r_[w['segments'][0]['points'][0],w['segments'][-1]['points'][-1]]
        np.testing.assert_allclose(row[2:8],ends,rtol=0,atol=1e-6)
        length=sum(np.linalg.norm(np.diff(np.array(s['points']),axis=0),axis=1).sum() for s in w['segments'])
        assert abs(float(cached[0][f'I{row_index}'])-length)<1e-5
        assert float(cached[0][f'J{row_index}'])==len(w['segments'])
    count=0
    for count,(row,e) in enumerate(zip(book['曲线点序列'].iter_rows(min_row=2,values_only=True),expected),1):
        wi,wid,si,pi,kind,q,step=e
        assert tuple(row[:4])==(wid,si,pi,kind)
        np.testing.assert_allclose(row[4:7],q,rtol=0,atol=1e-6)
        assert abs(float(cached[1][f'H{count+1}'])-step)<1e-7
    assert count==len(expected)==book['曲线点序列'].max_row-1
    book.close()
    csv.field_size_limit(10000000)
    with (out/'predicted_welds.csv').open(encoding='utf-8-sig',newline='') as f:rows=list(csv.DictReader(f))
    assert len(rows)==len(predictions)
    for r,w in zip(rows,predictions):
        assert r['weld_id']==w['id'] and r['type']==w['type']
        assert int(r['segment_count'])==len(w['segments'])
        assert json.loads(r['ordered_segments_json'])==w['segments']
        np.testing.assert_allclose([float(r[k]) for k in ['start_x_mm','start_y_mm','start_z_mm','end_x_mm','end_y_mm','end_z_mm']],
                                   np.r_[w['segments'][0]['points'][0],w['segments'][-1]['points'][-1]],rtol=0,atol=1e-9)
    with (out/'预测焊缝.ply').open('rb') as f:vertices=PlyData.read(f)['vertex']
    assert len(vertices)==len(expected)
    coords=np.column_stack([vertices[n] for n in ['x','y','z']])*1000.
    np.testing.assert_allclose(coords,np.array([e[5] for e in expected]),rtol=0,atol=1e-8)
    np.testing.assert_array_equal(vertices['weld_id'],[e[0] for e in expected])
    np.testing.assert_array_equal(vertices['segment_id'],[e[2] for e in expected])
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(out/'预测焊缝.xlsx') as z:
        assert z.testzip() is None
        styles=ET.fromstring(z.read('xl/styles.xml'))
        assert all(font.find('m:color',ns).attrib=={'rgb':'FF000000'} for font in styles.find('m:fonts',ns))
        assert all(x.attrib['fillId']=='0' for x in styles.find('m:cellXfs',ns))
        assert styles.find('m:fills',ns)[0].find('m:patternFill',ns).attrib['patternType']=='none'
    result={'welds':len(predictions),'curve_points':len(expected),'csv_excel_ply_all_coordinates_and_ids':'PASS',
            'black_white_no_fill':'PASS','scope':'File consistency only; extraction accuracy is in independent_score.json'}
    (out/'export_qa.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result),flush=True)


if __name__=='__main__':main()
