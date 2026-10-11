"""Offline inventory; does not assign train/test or feed labels to inference."""
import ast,collections,hashlib,json,pathlib,sys
import openpyxl
ROOT=pathlib.Path(__file__).resolve().parents[1]
DATA=ROOT.parent/'data/G_Data'
def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(4194304),b''):h.update(b)
 return h.hexdigest()
rows=[]
for i,p in enumerate(sorted(DATA.rglob('*_GW1_AllFaces.pcd'))):
 d=p.parent; record=dict(directory=str(d.relative_to(DATA)),pcd_sha256=sha(p),pcd_bytes=p.stat().st_size)
 with p.open('rb') as f:
  header=[]
  for _ in range(40):
   line=f.readline().decode('ascii').strip();header.append(line)
   if line.startswith('DATA '):break
 record['pcd_header']=header
 books=[]
 for b in sorted(d.glob('*.xlsx')):
  w=openpyxl.load_workbook(b,read_only=True,data_only=False)
  sheet=w['焊缝信息'] if '焊缝信息' in w.sheetnames else w.active
  it=sheet.iter_rows(values_only=True);headers=[str(x or '').split('\n')[-1] for x in next(it)]
  data=[dict(zip(headers,r)) for r in it if any(x is not None for x in r)]
  data=[r for r in data if r.get('ObjectId')]
  counts=collections.Counter(str(r.get('WeldType (isVerticalWeldLine)')) for r in data)
  books.append(dict(file=b.name,sha256=sha(b),parents=len(data),children=sum(int(r.get('SegmentCount') or 0) for r in data),types=dict(counts),curved_parents=sum(bool(r.get('HasMid (Any)')) for r in data),headers=headers))
  w.close()
 record['workbooks']=books;rows.append(record)
 if (i+1)%25==0:print('AUDITED',i+1,flush=True)
groups=collections.defaultdict(list)
for r in rows:groups[r['pcd_sha256']].append(r['directory'])
result=dict(cases=rows,pcd_byte_unique=len(groups),duplicate_groups=[v for v in groups.values() if len(v)>1],split_assigned=False,rigid_or_resampled_duplicates_not_yet_checked=True)
out=ROOT/'migration_checks';(out/'dataset_audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
missing=collections.defaultdict(set)
for p in (ROOT/'weld_pipeline').glob('*.py'):
 try:tree=ast.parse(p.read_text(encoding='utf-8-sig'))
 except (SyntaxError,UnicodeError):continue
 for n in ast.walk(tree):
  if isinstance(n,ast.ImportFrom) and n.module and n.module.split('.')[0] in ['training','weld_pipeline'] and n.level==0:
   dest=ROOT.joinpath(*n.module.split('.'))
   if not dest.with_suffix('.py').exists() and not dest.is_dir():missing[n.module].add(p.name)
(out/'missing_internal_modules.json').write_text(json.dumps({k:sorted(v) for k,v in missing.items()},indent=2))
print(json.dumps(dict(cases=len(rows),byte_unique=len(groups),duplicate_groups=len(result['duplicate_groups']),missing_internal_modules=len(missing)),ensure_ascii=False),flush=True)
