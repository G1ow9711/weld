"""Independent saved-output evaluation and export readback for the fit trial."""
import csv,json
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from plyfile import PlyData
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from weld_pipeline.truth import load_truth
from weld_pipeline.evaluation import score_welds
from training.labels import sha256


def main():
    run=Path('weld_pipeline/runs/contact_single_piece_v2');out=run/'raw_input_inference'
    prediction=json.loads((out/'predictions.json').read_text());truth,_=load_truth(Path('DataSet/9月20日/161-LB18A'))
    score=score_welds(prediction,truth);internal=[]
    for row in score['per_truth']:
        if not row['accepted']:continue
        ps=prediction[row['prediction_index']]['segments'];ts=truth[row['truth_index']]['segments']
        if row['reversed']:ps=[{**s,'points':s['points'][::-1]} for s in ps[::-1]]
        error=max(np.linalg.norm(np.asarray(p['points'])[[0,-1]]-np.asarray(t['points'])[[0,-1]],axis=1).max() for p,t in zip(ps,ts))
        internal.append(float(error));row['max_all_segment_endpoint_error_mm']=float(error)
    score['all_segment_endpoint_max_among_accepted_mm']=max(internal)
    score['accepted_welds_with_any_internal_endpoint_over_3mm']=sum(x>3 for x in internal)
    score['scope']='Training-assembly fitting diagnostic, not unseen-workpiece or formal acceptance.'
    (out/'independent_saved_output_score.json').write_text(json.dumps(score,indent=2),encoding='utf-8')
    book=load_workbook(out/'预测焊缝.xlsx',read_only=True,data_only=True)
    rows=list(book['焊缝汇总'].values);point_count=book['曲线点序列'].max_row-1
    assert len(rows)-1==len(prediction)
    for row,p in zip(rows[1:],prediction):
        assert row[0]==p['id']
        assert np.max(np.abs(np.array(row[2:8])-np.r_[p['segments'][0]['points'][0],p['segments'][-1]['points'][-1]]))<1e-6
        length=sum(np.linalg.norm(np.diff(np.array(s['points']),axis=0),axis=1).sum() for s in p['segments'])
        assert abs(row[8]-length)<1e-5
        assert row[9]==len(p['segments'])
    book.close()
    with (out/'预测焊缝.ply').open('rb') as f:ply=PlyData.read(f)
    assert len(ply['vertex'])==point_count
    with (out/'predicted_welds.csv').open(encoding='utf-8-sig',newline='') as f:
        csv.field_size_limit(10000000);csv_rows=list(csv.DictReader(f))
    assert len(csv_rows)==len(prediction)
    book=load_workbook(out/'预测焊缝.xlsx',read_only=False,data_only=False)
    assert all(c.fill.patternType is None for s in book for row in s for c in row)
    assert all(c.font.color and c.font.color.type=='rgb' and c.font.color.rgb=='FF000000' for s in book for row in s for c in row)
    book.close()
    qa={'welds':len(prediction),'curve_points':point_count,'csv_excel_ply_readback':'PASS','black_white_no_fill':'PASS',
        'checkpoint_sha256':sha256(run/'best.pt'),'scope':'diagnostic export validation only',
        'score':{k:v for k,v in score.items() if k not in ['per_truth','unmatched_prediction_indices']}}
    (out/'qa.json').write_text(json.dumps(qa,indent=2),encoding='utf-8')
    fig,axs=plt.subplots(1,2,figsize=(13,8),sharex=True,sharey=True)
    for ax,ws,title,color in zip(axs,[truth,prediction],['Reference: training assembly','Learned predictions: SAME training assembly'],['green','red']):
        for w in ws:
            for s in w['segments']:
                p=np.asarray(s['points']);ax.plot(p[:,0],p[:,1],c=color,lw=.7);ax.scatter(p[[0,-1],0],p[[0,-1],1],s=2,c='black')
        ax.set_aspect('equal');ax.set_title(title);ax.set_xlabel('X (mm)');ax.set_ylabel('Y (mm)')
    fig.tight_layout();fig.savefig(out/'training_piece_comparison.png',dpi=150);plt.close(fig)
    print(json.dumps(qa))


if __name__=='__main__':main()
